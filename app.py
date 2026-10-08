"""
Tabhi Supplier Outreach Agent (demo)
------------------------------------
A small local web app that drafts supplier emails from Tabhi's built-in
templates (New / Extend / Finish), lets you edit them, and sends them
directly over SMTP from a single fixed sender: nnargund@mondee.com.

No database: templates live in templates.json and every send attempt is
appended to logs/email_log.jsonl.

Run:  python app.py   then open http://localhost:5050
"""
import csv
import io
import json
import os
import re
import smtplib
import ssl
import threading
import time
import uuid
from datetime import datetime, timezone
from email.message import EmailMessage
from email.utils import formataddr, make_msgid, formatdate

from flask import Flask, jsonify, request, send_file, Response

BASE = os.path.dirname(os.path.abspath(__file__))
TEMPLATES_FILE = os.path.join(BASE, "templates.json")
LOG_DIR = os.path.join(BASE, "logs")
LOG_FILE = os.path.join(LOG_DIR, "email_log.jsonl")
CONFIG_FILE = os.path.join(BASE, "config.env")

# The only address this tool will ever send from.
SENDER_EMAIL = "nnargund@mondee.com"

os.makedirs(LOG_DIR, exist_ok=True)
_log_lock = threading.Lock()


# --------------------------------------------------------------------------
# Config (config.env, simple KEY=VALUE lines; environment variables win)
# --------------------------------------------------------------------------
def load_config():
    cfg = {
        "SMTP_HOST": "smtp.gmail.com",
        "SMTP_PORT": "587",
        "SMTP_SECURITY": "starttls",  # starttls | ssl | none
        "SMTP_USERNAME": SENDER_EMAIL,
        "SMTP_PASSWORD": "",
        "SENDER_NAME": "Nish | Tabhi",
    }
    if os.path.exists(CONFIG_FILE):
        raw = open(CONFIG_FILE, "rb").read()
        text = raw.decode("utf-16") if raw[:2] in (b"\xff\xfe", b"\xfe\xff") else raw.decode("utf-8-sig", "replace")
        if True:
            for line in text.splitlines():
                line = line.strip()
                if not line or line.startswith("#") or "=" not in line:
                    continue
                k, v = line.split("=", 1)
                cfg[k.strip()] = v.strip().strip('"').strip("'")
    for k in list(cfg):
        if os.environ.get(k):
            cfg[k] = os.environ[k]
    # Google shows App Passwords as "abcd efgh ijkl mnop"; the spaces aren't part of it
    cfg["SMTP_PASSWORD"] = cfg.get("SMTP_PASSWORD", "").replace(" ", "").replace("\u00a0", "")
    return cfg


def smtp_status(cfg):
    problems = []
    if not cfg.get("SMTP_PASSWORD"):
        problems.append("SMTP_PASSWORD is empty in config.env")
    if cfg.get("SMTP_USERNAME", "").lower() != SENDER_EMAIL:
        problems.append(f"SMTP_USERNAME must be {SENDER_EMAIL}")
    return problems


# --------------------------------------------------------------------------
# Templates (re-drafted from the GTM tool's built-in scripts)
# --------------------------------------------------------------------------
DEFAULT_TEMPLATES = {
    "about": (
        "Tabhi is a travel marketplace group with $3 billion in gross sales, operating across the US, "
        "Canada, Mexico and Brazil through our group companies, including Interep, Consolid and Orinter. "
        "Together we reach 125 million consumers worldwide and 65,000 travel agencies through our B2B business."
    ),
    "signature": "Best regards,\nNish\nTabhi\nnnargund@mondee.com",
    "types": {
        "New": {
            "label": "New supplier",
            "subject": "Partnership with Tabhi: {supplier} {category} for Miraee and MondeeONE",
            "body": (
                "{greeting}\n\n"
                "{about}\n\n"
                "We're now expanding our supplier partnerships directly to launch our Agentic AI Corporate Travel "
                "platform, Miraee, and our Community & Membership Org platform, MondeeONE, and {supplier} is a "
                "priority partner for {category} across {markets}.\n\n"
                "We'd like to explore an API partnership. Could we set up a 30-minute call in the next week to cover:\n"
                "- API access and sandbox credentials\n"
                "- Commercial model and terms\n"
                "- Which markets and currencies the agreement would cover\n"
                "- Typical certification and go-live timeline\n\n"
                "Looking forward to hearing from you.\n\n"
                "{signature}"
            ),
        },
        "Extend": {
            "label": "Extend an existing integration",
            "subject": "Extending our {supplier} integration to new markets",
            "body": (
                "{greeting}\n\n"
                "{about}\n\n"
                "As we prepare to launch our Agentic AI Corporate Travel platform, Miraee, and our Community & "
                "Membership Org platform, MondeeONE, we're already working with {supplier} and would like to extend "
                "our integration to more points of sale: {markets}.\n\n"
                "Could you let us know what's needed on your side (contract addendum, new credentials or office IDs, "
                "and any certification), and the likely timeline? A short call this week would help us plan.\n\n"
                "{signature}"
            ),
        },
        "Finish": {
            "label": "Finish an integration in progress",
            "subject": "Next steps on our {supplier} integration",
            "body": (
                "{greeting}\n\n"
                "I'm following up from Tabhi on our {supplier} integration, which is currently at the {stage} stage "
                "on our side. With $3 billion in gross sales, 125 million consumers and 65,000 travel agencies on our "
                "platforms, and the upcoming launch of Miraee, our Agentic AI Corporate Travel platform, and "
                "MondeeONE, our Community & Membership Org platform, we're keen to get it live. We'd like to agree on "
                "the remaining steps and target go-live in the next two weeks.\n\n"
                "Could we set up a 30-minute call to close out the open items?\n\n"
                "{signature}"
            ),
        },
    },
}

TYPES = ["New", "Extend", "Finish"]
FALLBACK = {"category": "travel content", "markets": "our core markets", "stage": "integration"}


def load_templates():
    if os.path.exists(TEMPLATES_FILE):
        try:
            with open(TEMPLATES_FILE, encoding="utf-8") as f:
                t = json.load(f)
            # fill anything missing from defaults
            merged = json.loads(json.dumps(DEFAULT_TEMPLATES))
            merged["about"] = t.get("about", merged["about"])
            merged["signature"] = t.get("signature", merged["signature"])
            for k in TYPES:
                if k in t.get("types", {}):
                    merged["types"][k].update({x: t["types"][k][x] for x in ("subject", "body") if x in t["types"][k]})
            return merged
        except Exception:
            pass
    return json.loads(json.dumps(DEFAULT_TEMPLATES))


def save_templates(t):
    with open(TEMPLATES_FILE, "w", encoding="utf-8") as f:
        json.dump(t, f, indent=2, ensure_ascii=False)


PLACEHOLDER = re.compile(r"\{(greeting|supplier|category|markets|stage|contact_name|about|signature)\}")


def render(text, values):
    out = PLACEHOLDER.sub(lambda m: values.get(m.group(1), ""), text)
    out = re.sub(r"[ \t]+\n", "\n", out)
    out = re.sub(r"\n{3,}", "\n\n", out)
    return out.strip()


def draft_for(row, tpl):
    kind = row.get("type") if row.get("type") in TYPES else "New"
    contact = (row.get("contact_name") or "").strip()
    first = contact.split()[0] if contact else ""
    values = {
        "supplier": row.get("supplier", "").strip(),
        "category": (row.get("category") or "").strip() or FALLBACK["category"],
        "markets": (row.get("markets") or "").strip() or FALLBACK["markets"],
        "stage": (row.get("stage") or "").strip() or FALLBACK["stage"],
        "contact_name": contact,
        "greeting": f"Hi {first}," if first else f"Hi {row.get('supplier', '').strip()} team,",
        "about": tpl["about"],
        "signature": tpl["signature"],
    }
    t = tpl["types"][kind]
    return {
        "type": kind,
        "subject": render(t["subject"], values),
        "body": render(t["body"], values),
    }


# --------------------------------------------------------------------------
# Spreadsheet parsing
# --------------------------------------------------------------------------
FIELDS = ["supplier", "email", "type", "contact_name", "category", "markets", "stage", "cc", "owner", "notes"]
ALIASES = {
    "supplier": ["suppliername", "supplier", "company", "companyname", "partner", "partnername", "vendor", "name"],
    "email": ["email", "emailid", "emailaddress", "toemail", "to", "mail", "contactemail", "emailids"],
    "type": ["type", "kind", "emailtype", "outreachtype", "template"],
    "contact_name": ["contactname", "contactperson", "contact", "firstname", "person", "recipientname"],
    "category": ["category", "categories", "product", "products", "productcategory", "content"],
    "markets": ["markets", "market", "pointsofsale", "pos", "regions", "region", "countries", "country"],
    "stage": ["stage", "currentstage", "status", "integrationstage"],
    "cc": ["cc", "ccemails", "ccemail", "copy"],
    "owner": ["owner", "supplierowner", "accountowner"],
    "notes": ["notes", "note", "comments", "comment", "remarks"],
}
EMAIL_RE = re.compile(r"^[^@\s,;]+@[^@\s,;]+\.[A-Za-z]{2,}$")


def norm(h):
    return re.sub(r"[^a-z0-9]", "", str(h or "").lower())


def map_headers(headers):
    mapping = {}
    used = set()
    normed = [norm(h) for h in headers]
    # exact alias matches, in alias priority order
    for field in FIELDS:
        for alias in ALIASES[field]:
            for i, h in enumerate(normed):
                if i not in used and h == alias and field not in mapping:
                    mapping[field] = i
                    used.add(i)
    return mapping


def split_emails(s):
    return [e.strip() for e in re.split(r"[,;\s]+", str(s or "")) if e.strip()]


def normalise_type(v):
    v = str(v or "").strip().lower()
    for t in TYPES:
        if v.startswith(t.lower()):
            return t
    return ""


def clean_row(r):
    row = {k: ("" if r.get(k) is None else str(r.get(k)).strip()) for k in FIELDS}
    issues = []
    if not row["supplier"]:
        issues.append("Supplier name missing")
    tos = split_emails(row["email"])
    if not tos:
        issues.append("Email missing")
    bad = [e for e in tos if not EMAIL_RE.match(e)]
    if bad:
        issues.append("Invalid email: " + ", ".join(bad))
    row["email"] = ", ".join(tos)
    ccs = split_emails(row["cc"])
    badcc = [e for e in ccs if not EMAIL_RE.match(e)]
    if badcc:
        issues.append("Invalid Cc: " + ", ".join(badcc))
    row["cc"] = ", ".join(ccs)
    t = normalise_type(row["type"])
    warnings = []
    if row["type"] and not t:
        warnings.append(f'Unknown type "{row["type"]}", using New')
    row["type"] = t or "New"
    if row["type"] == "Finish" and not row["stage"]:
        warnings.append("Finish email without a stage")
    if row["type"] in ("New", "Extend") and not row["markets"]:
        warnings.append("No markets given")
    if row["type"] == "New" and not row["category"]:
        warnings.append("No category given")
    row["issues"] = issues
    row["warnings"] = warnings
    row["id"] = uuid.uuid4().hex[:10]
    return row


def read_table(filename, data):
    name = filename.lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = [[("" if c is None else c) for c in r] for r in ws.iter_rows(values_only=True)]
    elif name.endswith(".csv") or name.endswith(".txt"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
    else:
        raise ValueError("Please upload a .xlsx or .csv file")
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        raise ValueError("The file is empty")
    # header = first row that maps both supplier and email (look at first 10 rows)
    for hi in range(min(10, len(rows))):
        m = map_headers(rows[hi])
        if "supplier" in m and "email" in m:
            break
    else:
        raise ValueError("Couldn't find the header row. The sheet needs at least 'Supplier name' and 'Email' columns.")
    out = []
    for r in rows[hi + 1:]:
        rec = {f: (r[i] if i < len(r) else "") for f, i in m.items()}
        out.append(clean_row(rec))
    ignored = [str(h) for i, h in enumerate(rows[hi]) if i not in m.values() and str(h).strip()]
    return out, sorted(m.keys(), key=FIELDS.index), ignored


# --------------------------------------------------------------------------
# Logging
# --------------------------------------------------------------------------
def append_log(entry):
    with _log_lock:
        with open(LOG_FILE, "a", encoding="utf-8") as f:
            f.write(json.dumps(entry, ensure_ascii=False) + "\n")


def read_logs():
    if not os.path.exists(LOG_FILE):
        return []
    out = []
    with open(LOG_FILE, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                try:
                    out.append(json.loads(line))
                except json.JSONDecodeError:
                    pass
    return out[::-1]  # newest first


# --------------------------------------------------------------------------
# Sending
# --------------------------------------------------------------------------
def smtp_connect(cfg):
    host, port, sec = cfg["SMTP_HOST"], int(cfg["SMTP_PORT"]), cfg.get("SMTP_SECURITY", "starttls").lower()
    ctx = ssl.create_default_context()
    if sec == "ssl":
        s = smtplib.SMTP_SSL(host, port, context=ctx, timeout=30)
    else:
        s = smtplib.SMTP(host, port, timeout=30)
        s.ehlo()
        if sec == "starttls":
            s.starttls(context=ctx)
            s.ehlo()
    if cfg.get("SMTP_PASSWORD"):
        s.login(cfg["SMTP_USERNAME"], cfg["SMTP_PASSWORD"])
    return s


def friendly_error(e):
    if isinstance(e, smtplib.SMTPAuthenticationError):
        return ("Login to the mail server failed for " + SENDER_EMAIL +
                ". Check the app password in config.env (it must be an App Password, not your normal password).")
    if isinstance(e, (TimeoutError, OSError)) and not isinstance(e, smtplib.SMTPException):
        return f"Couldn't reach the mail server ({e}). Check SMTP_HOST/SMTP_PORT and your network."
    return f"Mail server error: {e}"


def build_message(cfg, to_list, cc_list, subject, body):
    msg = EmailMessage()
    msg["From"] = formataddr((cfg.get("SENDER_NAME") or "", SENDER_EMAIL))
    msg["To"] = ", ".join(to_list)
    if cc_list:
        msg["Cc"] = ", ".join(cc_list)
    msg["Reply-To"] = SENDER_EMAIL
    msg["Subject"] = subject
    msg["Date"] = formatdate(localtime=True)
    msg["Message-ID"] = make_msgid(domain="mondee.com")
    msg.set_content(body)
    return msg


# --------------------------------------------------------------------------
# Web app
# --------------------------------------------------------------------------
app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024


@app.get("/")
def index():
    return send_file(os.path.join(BASE, "templates", "index.html"))


@app.get("/api/config")
def api_config():
    cfg = load_config()
    return jsonify({
        "sender": SENDER_EMAIL,
        "sender_name": cfg.get("SENDER_NAME", ""),
        "smtp_host": cfg["SMTP_HOST"],
        "smtp_port": cfg["SMTP_PORT"],
        "smtp_security": cfg.get("SMTP_SECURITY"),
        "problems": smtp_status(cfg),
        "fallbacks": FALLBACK,
    })


@app.get("/api/templates")
def api_templates():
    return jsonify(load_templates())


@app.post("/api/templates")
def api_templates_save():
    data = request.get_json(force=True)
    t = load_templates()
    t["about"] = data.get("about", t["about"])
    t["signature"] = data.get("signature", t["signature"])
    for k in TYPES:
        if k in data.get("types", {}):
            for x in ("subject", "body"):
                if x in data["types"][k]:
                    t["types"][k][x] = data["types"][k][x]
    save_templates(t)
    return jsonify(t)


@app.post("/api/templates/reset")
def api_templates_reset():
    if os.path.exists(TEMPLATES_FILE):
        os.remove(TEMPLATES_FILE)
    return jsonify(load_templates())


@app.post("/api/parse")
def api_parse():
    f = request.files.get("file")
    if not f:
        return jsonify({"error": "No file uploaded"}), 400
    try:
        rows, mapped, ignored = read_table(f.filename, f.read())
    except Exception as e:
        return jsonify({"error": str(e)}), 400
    # duplicates (same supplier + same email)
    seen = {}
    for r in rows:
        key = (r["supplier"].lower(), r["email"].lower())
        if key in seen and r["supplier"]:
            r["warnings"].append(f"Duplicate of row {seen[key]}")
        else:
            seen[key] = rows.index(r) + 1
    return jsonify({"rows": rows, "mapped": mapped, "ignored": ignored, "filename": f.filename})


@app.post("/api/validate")
def api_validate():
    rows = request.get_json(force=True).get("rows", [])
    return jsonify({"rows": [clean_row(r) for r in rows]})


@app.post("/api/draft")
def api_draft():
    rows = request.get_json(force=True).get("rows", [])
    tpl = load_templates()
    out = []
    for r in rows:
        d = draft_for(r, tpl)
        out.append({**r, **d, "to": r.get("email", ""), "cc": r.get("cc", "")})
    return jsonify({"drafts": out})


@app.post("/api/send")
def api_send():
    data = request.get_json(force=True)
    drafts = data.get("drafts", [])
    redirect = bool(data.get("demo_redirect", True))
    cfg = load_config()
    problems = smtp_status(cfg)
    batch = uuid.uuid4().hex[:8]
    results = []
    if problems:
        return jsonify({"error": "Email isn't set up yet: " + "; ".join(problems)}), 400
    try:
        conn = smtp_connect(cfg)
    except Exception as e:
        err = friendly_error(e)
        for d in drafts:  # log the failed attempt for every email in the batch
            append_log({
                "time": datetime.now(timezone.utc).isoformat(timespec="seconds"), "batch": batch,
                "supplier": d.get("supplier", ""), "contact_name": d.get("contact_name", ""),
                "type": d.get("type", "New"), "from": SENDER_EMAIL, "intended_to": d.get("to", ""),
                "intended_cc": d.get("cc", ""), "to": "", "cc": "", "subject": d.get("subject", ""),
                "body": d.get("body", ""), "edited": bool(d.get("edited")), "demo_redirect": redirect,
                "owner": d.get("owner", ""), "source": d.get("source", ""), "status": "Failed",
                "error": err, "message_id": "",
            })
        return jsonify({"error": err}), 502
    try:
        for d in drafts:
            to_list = split_emails(d.get("to"))
            cc_list = split_emails(d.get("cc"))
            subject = (d.get("subject") or "").strip()
            body = d.get("body") or ""
            entry = {
                "time": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                "batch": batch,
                "supplier": d.get("supplier", ""),
                "contact_name": d.get("contact_name", ""),
                "type": d.get("type", "New"),
                "from": SENDER_EMAIL,
                "intended_to": ", ".join(to_list),
                "intended_cc": ", ".join(cc_list),
                "subject": subject,
                "body": body,
                "edited": bool(d.get("edited")),
                "demo_redirect": redirect,
                "owner": d.get("owner", ""),
                "source": d.get("source", ""),
            }
            bad = [e for e in to_list + cc_list if not EMAIL_RE.match(e)]
            if not to_list or bad or not subject:
                entry.update(status="Failed", error="Missing or invalid recipient/subject" + (": " + ", ".join(bad) if bad else ""),
                             to="", cc="", message_id="")
                append_log(entry)
                results.append({"id": d.get("id"), "status": "Failed", "error": entry["error"]})
                continue
            send_to, send_cc = (to_list, cc_list)
            if redirect:
                send_to, send_cc = [SENDER_EMAIL], []
                body = (f"[DEMO MODE: this email would have gone to {', '.join(to_list)}"
                        + (f", cc {', '.join(cc_list)}" if cc_list else "") + "]\n\n" + body)
            msg = build_message(cfg, send_to, send_cc, subject, body)
            entry.update(to=", ".join(send_to), cc=", ".join(send_cc), message_id=msg["Message-ID"])
            try:
                try:
                    conn.send_message(msg, from_addr=SENDER_EMAIL, to_addrs=send_to + send_cc)
                except smtplib.SMTPServerDisconnected:
                    conn = smtp_connect(cfg)
                    conn.send_message(msg, from_addr=SENDER_EMAIL, to_addrs=send_to + send_cc)
                entry.update(status="Sent", error="")
            except Exception as e:
                entry.update(status="Failed", error=friendly_error(e)[:300])
            append_log(entry)
            results.append({"id": d.get("id"), "status": entry["status"], "error": entry["error"],
                            "time": entry["time"]})
            time.sleep(0.4)  # be gentle with the mail server
    finally:
        try:
            conn.quit()
        except Exception:
            pass
    return jsonify({"batch": batch, "results": results})


@app.post("/api/test-email")
def api_test_email():
    cfg = load_config()
    problems = smtp_status(cfg)
    if problems:
        return jsonify({"ok": False, "error": "; ".join(problems)}), 400
    try:
        conn = smtp_connect(cfg)
        msg = build_message(cfg, [SENDER_EMAIL], [], "Tabhi outreach agent: test email",
                            "This is a test from the Tabhi supplier outreach agent. Sending works.")
        conn.send_message(msg, from_addr=SENDER_EMAIL, to_addrs=[SENDER_EMAIL])
        conn.quit()
        return jsonify({"ok": True})
    except Exception as e:
        return jsonify({"ok": False, "error": friendly_error(e)}), 502


@app.get("/api/logs")
def api_logs():
    return jsonify({"logs": read_logs()})


@app.get("/api/logs.csv")
def api_logs_csv():
    logs = read_logs()
    cols = ["time", "batch", "status", "type", "supplier", "contact_name", "from", "intended_to", "intended_cc",
            "to", "cc", "demo_redirect", "edited", "subject", "error", "message_id", "owner", "source", "body"]
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=cols, extrasaction="ignore")
    w.writeheader()
    for l in logs:
        w.writerow(l)
    return Response(buf.getvalue().encode("utf-8-sig"), mimetype="text/csv",
                    headers={"Content-Disposition": "attachment; filename=email_log.csv"})


@app.get("/api/sample.<ext>")
def api_sample(ext):
    headers = ["Supplier name", "Email", "Type", "Contact name", "Category", "Markets", "Stage", "Cc", "Owner", "Notes"]
    rows = [
        ["Trip.com", "partnerships@example.com", "New", "", "Air, hotel, packages", "UAE, India", "", "", "Shreepal & Sheriff", "Also majority owner of Travelfusion"],
        ["Sabre Cars", "cars@example.com", "Extend", "", "Car rental", "Mexico, Brazil", "", "", "Shreepal & Sheriff", ""],
        ["Klook", "klook@example.com", "Finish", "Priya Shah", "Activities & attractions", "", "Staging / UAT", "", "Shreepal & Sheriff", ""],
    ]
    if ext == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(headers)
        w.writerows(rows)
        return Response(buf.getvalue().encode("utf-8-sig"), mimetype="text/csv",
                        headers={"Content-Disposition": "attachment; filename=supplier_outreach_template.csv"})
    if ext == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Suppliers"
        ws.append(headers)
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="172031")
        for col, wdt in zip("ABCDEFGHIJ", [22, 30, 10, 16, 26, 22, 16, 26, 20, 36]):
            ws.column_dimensions[col].width = wdt
        ws.freeze_panes = "A2"
        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)
        return send_file(bio, as_attachment=True, download_name="supplier_outreach_template.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    return "Not found", 404


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "5050"))
    print(f"\n  Tabhi outreach agent running at  http://localhost:{port}\n  Sender: {SENDER_EMAIL}\n")
    app.run(host="127.0.0.1", port=port, debug=False)
